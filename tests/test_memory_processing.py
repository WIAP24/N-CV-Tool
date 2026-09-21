import builtins
import io
import json
import sys
import unittest
from contextlib import ExitStack
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch
from zipfile import ZipFile

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from niras_cv_screener.cloud_files import MemoryCV
from niras_cv_screener.workflow import process_paths
from niras_cv_screener.llm import screen_cv_with_metadata
from test_criteria_and_scoring import make_result, sample_criteria


class MemoryProcessingTests(unittest.TestCase):
    def process_without_writes(self, failed=False):
        original_open = builtins.open
        original_io_open = io.open

        def guarded(opener):
            def check(file, mode="r", *args, **kwargs):
                if any(flag in str(mode) for flag in "wax+"):
                    raise AssertionError("Disk write attempted")
                return opener(file, mode, *args, **kwargs)
            return check

        def assessment(**kwargs):
            if failed:
                raise RuntimeError("private CV contents in provider error")
            return {"result": make_result("Synthetic Candidate", kwargs["file_name"], 4, 4),
                    "usage": {"input_tokens": 100, "output_tokens": 100}}

        with ExitStack() as stack:
            stack.enter_context(patch("builtins.open", side_effect=guarded(original_open)))
            stack.enter_context(patch("io.open", side_effect=guarded(original_io_open)))
            for target in ("tempfile.NamedTemporaryFile", "tempfile.TemporaryDirectory",
                           "tempfile.mkstemp", "tempfile.mkdtemp",
                           "openpyxl.worksheet._writer.create_temporary_file", "pathlib.Path.mkdir"):
                stack.enter_context(patch(target, side_effect=AssertionError("Temporary disk storage attempted")))
            model = stack.enter_context(patch("niras_cv_screener.workflow.screen_cv_with_metadata", side_effect=assessment))
            result = process_paths(
                [MemoryCV("synthetic.txt", b"Synthetic experience and qualifications. " * 30)],
                sample_criteria(), "test-only-key", "gpt-4o-mini", output_root=None,
                use_result_cache=True, enable_model_comparison=True, comparison_model="gpt-4o-mini",
            )
            self.assertEqual(model.call_count, 1 if failed else 2)
        return result

    def test_full_reports_without_disk_writes(self):
        result = self.process_without_writes()
        self.assertEqual(result["processed"], 1)
        self.assertEqual(result["skipped"], 0)
        workbook = load_workbook(io.BytesIO(result["excel_download"]))
        self.assertIn("Summary", workbook.sheetnames)
        self.assertIn("Synthetic Candidate", [cell.value for row in workbook["Summary"] for cell in row])
        with ZipFile(io.BytesIO(result["zip_download"])) as archive:
            self.assertIn("results.json", archive.namelist())
            self.assertIn("extracted_text/synthetic.txt", archive.namelist())
            self.assertTrue(any(name.startswith("comparison_results/") for name in archive.namelist()))
            manifest = json.loads(archive.read("run_manifest.json"))
            self.assertEqual(manifest["storage_mode"], "session_memory")
            self.assertEqual(manifest["extraction_cache"], "disabled")
            self.assertFalse(manifest["use_result_cache"])
            self.assertNotIn(b"test-only-key", b"".join(archive.read(name) for name in archive.namelist()))

    def test_failed_run_reports_without_raw_provider_error(self):
        result = self.process_without_writes(failed=True)
        self.assertEqual(result["processed"], 0)
        self.assertEqual(result["skipped"], 1)
        self.assertNotIn("private CV contents", str(result["errors"]))
        with ZipFile(io.BytesIO(result["zip_download"])) as archive:
            self.assertEqual(len(json.loads(archive.read("skipped_files.json"))), 1)

    def test_memory_mode_rejects_ocr(self):
        with self.assertRaisesRegex(ValueError, "OCR is disabled"):
            process_paths([MemoryCV("synthetic.pdf", b"test")], sample_criteria(),
                          "test-only", "gpt-4o-mini", None, use_ocr=True)

    def test_api_storage_disabled_including_retries(self):
        client = Mock()
        response = SimpleNamespace(output_text=json.dumps(make_result("Synthetic", "example.txt", 4, 4)), usage={})
        client.responses.create.side_effect = [TypeError("seed unsupported"), response]
        with patch("niras_cv_screener.llm.openai_client", return_value=client):
            screen_cv_with_metadata(sample_criteria(), "synthetic content", "example.txt", "test", "gpt-4o-mini")
        self.assertEqual(client.responses.create.call_count, 2)
        for call in client.responses.create.call_args_list:
            self.assertIs(call.kwargs["store"], False)

    def test_gpt5_mini_request_parameters(self):
        client = Mock()
        client.responses.create.return_value = SimpleNamespace(
            output_text=json.dumps(make_result("Synthetic", "example.txt", 4, 4)), usage={})
        with patch("niras_cv_screener.llm.openai_client", return_value=client):
            screen_cv_with_metadata(sample_criteria(), "synthetic", "example.txt", "test", "gpt-5-mini", "medium")
        request = client.responses.create.call_args.kwargs
        self.assertNotIn("temperature", request)
        self.assertNotIn("seed", request)
        self.assertEqual(request["reasoning"], {"effort": "medium"})
        self.assertIs(request["store"], False)


if __name__ == "__main__":
    unittest.main()
